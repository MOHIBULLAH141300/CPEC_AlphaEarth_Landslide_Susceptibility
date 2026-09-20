from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(
    r"D:\DING PROJECT\00_READ_ME_FIRST_2018_V3_RESULTS"
    r"\CPEC_KKH_updated_next_methodology_summary_for_teacher_2026-05-10.docx"
)
REPORT_COPY = Path(
    r"D:\DING PROJECT\05_reports"
    r"\CPEC_KKH_updated_next_methodology_summary_for_teacher_2026-05-10.docx"
)

ACCENT = "1F4E79"
LIGHT = "EAF2F8"
MID = "D9EAF7"
TEXT = RGBColor(33, 33, 33)
MUTED = RGBColor(90, 90, 90)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def table_geometry(table, widths: list[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths)))

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_field(paragraph, field_code: str) -> None:
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = field_code
    fld_char_sep = OxmlElement("w:fldChar")
    fld_char_sep.set(qn("w:fldCharType"), "separate")
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_sep)
    run._r.append(fld_char_end)


def add_bottom_border(paragraph, color="B7C9D6", size="8") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.3)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(10.6)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(5)

    for name, size in [("Title", 22), ("Subtitle", 11), ("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11)]:
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(31, 78, 121) if name != "Subtitle" else MUTED
        if name.startswith("Heading") or name == "Title":
            style.font.bold = True
        style.paragraph_format.space_before = Pt(9 if name == "Heading 1" else 6)
        style.paragraph_format.space_after = Pt(4)

    for list_style in ["List Bullet", "List Number"]:
        style = styles[list_style]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(10.4)
        style.paragraph_format.space_after = Pt(4)


def add_header_footer(doc: Document) -> None:
    section = doc.sections[0]
    header = section.header
    p = header.paragraphs[0]
    p.text = "CPEC/KKH Landslide Susceptibility - Updated Next Methodology"
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.runs[0].font.name = "Arial"
    p.runs[0].font.size = Pt(8.5)
    p.runs[0].font.color.rgb = MUTED
    add_bottom_border(p)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = fp.add_run("Page ")
    run.font.name = "Arial"
    run.font.size = Pt(8.5)
    run.font.color.rgb = MUTED
    add_field(fp, "PAGE")


def add_meta_table(doc: Document) -> None:
    data = [
        ("Prepared for", "Supervisor / Teacher"),
        ("Project", "CPEC/KKH landslide susceptibility mapping"),
        ("Date", "10 May 2026"),
        ("Purpose", "Updated summary of the next complete methodology after reviewing teacher-sent recent literature and current 2018 results"),
    ]
    table = doc.add_table(rows=len(data), cols=2)
    table.style = "Table Grid"
    table_geometry(table, [2100, 7600])
    for i, (label, value) in enumerate(data):
        table.cell(i, 0).text = label
        table.cell(i, 1).text = value
        set_cell_shading(table.cell(i, 0), LIGHT)
        for cell in table.rows[i].cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(10)
            table.cell(i, 0).paragraphs[0].runs[0].font.bold = True
    doc.add_paragraph()


def add_callout(doc: Document, title: str, body: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table_geometry(table, [9700])
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F4F8FB")
    p = cell.paragraphs[0]
    r = p.add_run(title + ": ")
    r.bold = True
    r.font.color.rgb = RGBColor(31, 78, 121)
    r.font.name = "Arial"
    r.font.size = Pt(10.6)
    r2 = p.add_run(body)
    r2.font.name = "Arial"
    r2.font.size = Pt(10.4)
    doc.add_paragraph()


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)


def add_simple_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table_geometry(table, widths)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = h
        set_cell_shading(cell, MID)
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.name = "Arial"
                run.font.size = Pt(9.4)
                run.font.color.rgb = RGBColor(31, 78, 121)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(9.2)
    doc.add_paragraph()


def build_doc() -> Document:
    doc = Document()
    style_document(doc)
    add_header_footer(doc)

    title = doc.add_paragraph(style="Title")
    title.add_run("Updated Next Methodology Summary")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run(
        "CPEC/KKH landslide susceptibility mapping after review of teacher-sent 2024-2025 literature"
    )
    add_meta_table(doc)

    add_callout(
        doc,
        "Main proposed direction",
        "Transferability-aware dynamic landslide susceptibility mapping of the CPEC/KKH corridor using AlphaEarth foundation embeddings and Spatial-CV stacked ensembles.",
    )

    doc.add_heading("1. Current Position Of The Project", level=1)
    doc.add_paragraph(
        "We have completed a corrected 2018 baseline for the official CPEC study area, including Pakistan and Kashgar, Xinjiang. "
        "This baseline should now be treated as the clean foundation for the next phase, not as a preliminary result."
    )
    add_simple_table(
        doc,
        ["Completed component", "Status", "Why it matters"],
        [
            ["Official study area", "Complete", "All modelling is clipped to the official CPEC boundary."],
            ["2018 corrected factor database", "Complete", "Includes terrain, hydro-climate, vegetation, land cover, geology, soil, seismicity, road distance, river distance, and fault distance."],
            ["Multicollinearity gate", "Complete", "Final conventional factors passed VIF screening; largest final VIF is below 4."],
            ["Model family", "Complete", "Logistic Regression, Random Forest, Extra Trees, XGBoost, LightGBM, CatBoost, and Spatial-CV Stacked Ensemble were evaluated."],
            ["Feature-set comparison", "Complete for 2018", "Conventional, AlphaEarth Embeddings, and Conventional + AlphaEarth Embeddings were compared."],
            ["Explainability", "Complete for 2018", "SHAP importance, beeswarm, and waterfall-style explanation outputs are available."],
            ["Uncertainty and exposure", "Complete for 2018", "Model-disagreement uncertainty, reliability, AlphaEarth added-value, and KKH/CPEC road hotspot outputs were created."],
        ],
        [2550, 1600, 5550],
    )

    doc.add_heading("2. What The Recent Literature Suggests", level=1)
    doc.add_paragraph(
        "The teacher-sent papers point toward transfer learning, domain adaptation, causal or counterfactual robustness, meta-learning, uncertainty, and stronger interpretation. "
        "For our data, the best strategy is to incorporate these ideas without changing the study into an image-segmentation project."
    )
    add_simple_table(
        doc,
        ["Literature direction", "How we will use it in this study"],
        [
            ["Transfer learning and data-scarce LSM", "Test whether the CPEC model transfers across different geomorphic and corridor subregions."],
            ["Multi-source domain transfer", "Treat CPEC as several source/target domains instead of one uniform region."],
            ["Cross-regional extrapolation", "Use leave-one-domain-out validation and explain where the model generalizes poorly."],
            ["Counterfactual / causal robustness", "Run ablation and scenario tests, especially for roads, rainfall, AlphaEarth, and seismic forcing."],
            ["Meta-learning", "Keep the Spatial-CV stacked ensemble as the framework model and report individual learner performance transparently."],
            ["XAI and uncertainty", "Use SHAP, calibration, repeated non-landslide sampling, and reliability maps to show trustworthiness."],
        ],
        [3100, 6600],
    )

    doc.add_heading("3. Proposed Complete Methodology", level=1)
    doc.add_paragraph(
        "The next methodology will be organized into eight connected modules. This structure keeps the work scientifically strong and easy for reviewers to follow."
    )
    add_simple_table(
        doc,
        ["Module", "Method", "Main outputs"],
        [
            ["1", "Inventory and factor database using the official CPEC boundary and corrected 2018 sample/factor table.", "Clean inventory, final factor list, study-area map."],
            ["2", "Multicollinearity and feature-stability gate using VIF, correlation checks, and optional permutation/RFE stability.", "Final selected factors and stability table."],
            ["3", "Spatial-CV model family and stacked ensemble for three feature sets: Conventional, AlphaEarth Embeddings, and Conventional + AlphaEarth Embeddings.", "AUC, PR-AUC, F1, balanced accuracy, Brier score, calibration, ROC/PR curves."],
            ["4", "Transferability and area-of-applicability testing across CPEC subdomains.", "Subdomain map, leave-one-domain-out metrics, domain-similarity table, transfer confidence map."],
            ["5", "Annual dynamic-factor susceptibility mapping for 2017-2024 using fixed 2018 trained models and annual predictors.", "Annual probability maps, year-to-year change maps, trend map, mean susceptibility, persistent hotspot map."],
            ["6", "AlphaEarth added-value and embedding-change analysis.", "Fused-minus-Conventional map, AlphaEarth disagreement map, annual embedding-change map, grouped SHAP contribution."],
            ["7", "Uncertainty, sampling robustness, and calibration.", "Repeated sampling uncertainty, reliable high-probability zones, uncertain high-probability zones, calibration curves."],
            ["8", "KKH/CPEC infrastructure exposure and scenario branch.", "Road exposure tables, persistent KKH hotspot segments, rainfall/PGA scenario maps, road-distance ablation result."],
        ],
        [900, 4750, 4050],
    )

    doc.add_heading("4. Feature Sets To Be Used", level=1)
    add_simple_table(
        doc,
        ["Feature set", "Predictors", "Purpose"],
        [
            ["Conventional", "19 physically interpretable factors: terrain, rainfall, vegetation, land cover, geology, soil, seismicity, road/river/fault distance.", "Baseline physically explainable susceptibility model."],
            ["AlphaEarth Embeddings", "64 annual AlphaEarth/Satellite Embedding bands from Google Earth Engine, available from 2017 onward.", "Foundation-model remote-sensing representation and annual context."],
            ["Conventional + AlphaEarth Embeddings", "All conventional factors plus the 64 embedding bands.", "Main fused model combining physical interpretability and remote-sensing representation."],
        ],
        [2500, 4550, 2650],
    )

    doc.add_heading("5. Dynamic Susceptibility Strategy", level=1)
    doc.add_paragraph(
        "For the annual extension, the model will not be retrained separately for every year. Instead, the corrected 2018 trained model will be held fixed and applied to annual predictor stacks from 2017 to 2024. "
        "This makes the year-to-year changes easier to interpret because model structure remains constant while annual rainfall, vegetation, land cover, and AlphaEarth embeddings change."
    )
    add_bullets(
        doc,
        [
            "Resolution: 250 m modelling grid to avoid unnecessary storage and false precision.",
            "Annual changing factors: monsoon rainfall, maximum 1-day rainfall, NDVI median, NDVI amplitude, land-cover class, and AlphaEarth embeddings.",
            "Fixed or semi-static factors: terrain, lithology, soil, distance to faults, distance to rivers/streams, distance to roads, and earthquake density until better annual seismic data are added.",
            "Outputs: annual susceptibility maps, previous-year difference maps, trend maps, mean susceptibility, persistent hotspot zones, and annual road exposure tables.",
        ],
    )

    doc.add_heading("6. Transferability Plan", level=1)
    doc.add_paragraph(
        "The major new scientific addition is transferability. CPEC is too diverse to treat as one uniform region, so we will test whether the model generalizes across different geomorphic and infrastructure domains."
    )
    add_numbered(
        doc,
        [
            "Create CPEC subdomains, such as KKH/Kashgar high mountains, northern Pakistan mountain belt, Indus basin, western Pakistan/Balochistan, and southern/coastal corridor.",
            "Run leave-one-domain-out validation: train on all domains except one and test on the held-out domain.",
            "Measure domain similarity using KL divergence, Jensen-Shannon divergence, or MMD across selected predictors.",
            "Create area-of-applicability or environmental dissimilarity maps to show where predictions are more reliable or less reliable.",
            "Compare Conventional, AlphaEarth Embeddings, and Conventional + AlphaEarth Embeddings under the same transferability tests.",
        ],
    )

    doc.add_heading("7. Uncertainty And Robustness Plan", level=1)
    add_simple_table(
        doc,
        ["Robustness check", "Reason", "Output"],
        [
            ["Repeated non-landslide sampling", "Absence/control-point choice can affect susceptibility maps.", "Mean probability, standard deviation, and confidence masks."],
            ["Calibration and Brier score", "A high AUC does not guarantee reliable probabilities.", "Calibration curves and Brier scores."],
            ["Feature-set disagreement", "Different predictor groups can disagree in uncertain zones.", "Disagreement and reliability maps."],
            ["Road-distance ablation", "Road distance is both a conditioning factor and exposure concern.", "Model and hotspot comparison with/without road distance."],
            ["Rainfall and seismic scenarios", "Teacher requested scenario thinking for extreme rainfall and earthquake effects.", "Scenario susceptibility maps and infrastructure exposure changes."],
        ],
        [2600, 4250, 2850],
    )

    doc.add_heading("8. Expected Final Outputs For The Teacher/Reviewers", level=1)
    add_bullets(
        doc,
        [
            "A publication-ready methodology flowchart.",
            "Final factor database, VIF table, and factor justification table.",
            "Spatial-CV model performance tables and ROC/PR/calibration curves.",
            "SHAP and grouped AlphaEarth interpretation figures.",
            "2018 baseline probability maps for all three feature sets.",
            "Annual dynamic susceptibility maps from 2017 to 2024.",
            "Annual change, trend, mean susceptibility, and persistent hotspot maps.",
            "Transferability and area-of-applicability maps.",
            "Sampling-robust uncertainty and reliability maps.",
            "KKH/CPEC road exposure hotspot tables and maps.",
            "Optional rainfall/PGA scenario maps if data availability is acceptable.",
        ],
    )

    doc.add_heading("9. Limitations And Careful Wording", level=1)
    add_bullets(
        doc,
        [
            "We should call the current result a 2018 dynamic-factor landslide susceptibility baseline.",
            "We should call the next maps annual dynamic-factor susceptibility maps for 2017-2024.",
            "We should avoid claiming fully validated temporal susceptibility until more dated landslide events are added.",
            "Deep CNN/GNN methods from some teacher-sent papers are mainly landslide detection/segmentation methods; we will borrow transfer/spatial-context ideas but keep the susceptibility framework suitable for our data.",
            "Counterfactual and ablation tests will be described as robustness diagnostics, not as formal causal proof.",
        ],
    )

    doc.add_heading("10. Immediate Next Steps", level=1)
    add_numbered(
        doc,
        [
            "Complete the annual 2017-2024 predictor exports and probability maps.",
            "Create the CPEC subdomain map for transferability testing.",
            "Run leave-one-domain-out model evaluation.",
            "Create area-of-applicability and domain-similarity outputs.",
            "Run repeated non-landslide sampling uncertainty experiments.",
            "Extend KKH/CPEC road exposure analysis to annual maps.",
            "Prepare the final methodology flowchart and manuscript-ready method section.",
        ],
    )

    doc.add_heading("Key References Supporting This Direction", level=1)
    refs = [
        "Singh et al. (2024), Scientific Reports: Ensembled transfer learning approach for error reduction in landslide susceptibility mapping of a data-scarce region. https://www.nature.com/articles/s41598-024-76541-4",
        "Halder et al. (2025), Scientific Reports: Improving landslide susceptibility prediction through ensemble recursive feature elimination and meta-learning. https://www.nature.com/articles/s41598-025-87587-3",
        "Zhao et al. (2025), GIScience & Remote Sensing: Counterfactual inference causal representation learning for fine-grained landslide mapping. https://www.tandfonline.com/doi/full/10.1080/15481603.2025.2598078",
        "Luo et al. (2025), Geoscience Frontiers: Landslide detection based on transfer learning and graph neural network. https://www.sciencedirect.com/science/article/pii/S1674987125001768",
        "Su et al. (2025), Geoscience Frontiers: Complex cross-regional LSM by multi-source domain transfer learning. https://www.sciencedirect.com/science/article/pii/S1674987125000581",
        "Wang et al. (2025), Geoscience Frontiers: Cross-regional extrapolation of landslide susceptibility mapping via transfer learning. https://www.sciencedirect.com/science/article/pii/S1674987125002178",
        "Google Earth Engine: Satellite Embedding V1 / AlphaEarth Foundations annual 64-band embeddings. https://developers.google.com/earth-engine/datasets/catalog/GOOGLE_SATELLITE_EMBEDDING_V1_ANNUAL",
    ]
    for ref in refs:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(ref)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Prepared as an updated methodology summary for supervisor review.")
    r.italic = True
    r.font.color.rgb = MUTED
    r.font.size = Pt(9.5)
    return doc


def main() -> None:
    doc = build_doc()
    OUT.parent.mkdir(parents=True, exist_ok=True)
    REPORT_COPY.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    doc.save(REPORT_COPY)
    print(OUT)
    print(REPORT_COPY)


if __name__ == "__main__":
    main()
