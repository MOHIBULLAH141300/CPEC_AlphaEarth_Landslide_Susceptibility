from pathlib import Path
import shutil

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


SOURCE = Path(r"C:\Users\Administrator\Desktop\cpec landslides\suggestions-renew.docx")
OUT_MAIN = Path(
    r"D:\DING PROJECT\05_reports\teacher_suggestions_filled_original_format_2026-05-11.docx"
)
OUT_COPY = Path(
    r"C:\Users\Administrator\Desktop\cpec landslides\suggestions-renew_filled_original_format_2026-05-11.docx"
)


def insert_paragraph_after(paragraph: Paragraph, text: str = "", style=None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style is not None:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def add_response_after(paragraph: Paragraph, lines: list[str]) -> Paragraph:
    anchor = paragraph
    for idx, line in enumerate(lines):
        p = insert_paragraph_after(anchor, style=paragraph.style)
        if idx == 0:
            run = p.add_run("【回复 / 执行方案】")
            run.bold = True
            p.add_run(line)
        else:
            p.add_run(line)
        anchor = p
    return anchor


def find_paragraph(doc: Document, starts_with: str) -> Paragraph:
    for p in doc.paragraphs:
        if p.text.strip().startswith(starts_with):
            return p
    raise ValueError(f"Cannot find paragraph starting with: {starts_with}")


RESPONSES = {
    "“Spatial-temporal evolution": [
        "拟采用升级后的题目，但将方法表达进一步具体化为：Spatial-temporal evolution of landslide susceptibility along the China-Pakistan Economic Corridor revealed by Spatial-CV stacked ensemble learning and AlphaEarth embeddings。",
        "中文理解：以官方CPEC研究区为对象，将2018年作为严格验证的基准年，并用2017-2024年度动态因子开展时序迁移/动态易发性压力测试。"
    ],
    "1. 充分阐明选用集成机器学习方法的Motivation": [
        "论文Motivation将按“研究区挑战—数据响应—模型响应”的逻辑重写。研究区挑战包括CPEC尺度大、跨越高山/丘陵/干旱高原/低地走廊、孕灾环境高度异质、样本空间不均衡、边界区域易出现外推预测。",
        "数据响应为：常规孕灾因子（地形、水文气候、植被/土地覆盖、岩性/土壤、地震、道路/河流/断层距离）+ AlphaEarth年度高维嵌入特征；模型响应为：Spatial-CV集成/堆叠模型、AOA可靠性约束、域间迁移验证和TreeSHAP解释。",
        "当前2018结果已对应此逻辑：官方边界、19个常规因子、AlphaEarth Embeddings、三类特征集对比、Spatial-CV Stacked Ensemble、ROC/PR/Calibration、SHAP、AOA/transfer confidence、道路暴露和域间对比均已完成。"
    ],
    "2. 阐明选用Extra Trees Classifier": [
        "方法部分将补充基础模型和元学习器的明确理由。Extra Trees通过随机切分和多树平均降低方差，对噪声、异常值和非线性阈值关系较稳健；XGBoost/LightGBM擅长拟合复杂交互和稀疏/高维特征；CatBoost对混合型变量和小中样本条件下的稳定性较好。",
        "本研究不只报告单个模型，而是让不同基础学习器在空间交叉验证框架下产生互补预测，再由二层Logistic meta-learner综合它们的证据。这样可以解释为什么集成模型适合CPEC这种样本不均衡、影响因子多、地貌差异明显的大尺度制图问题。",
        "写作时会避免不严谨表述：缺失值/NaN由统一预处理、掩膜和质量控制处理；树模型优势主要表述为对非线性、多因子交互、噪声和中小样本的鲁棒建模能力。"
    ],
    "3. 绘制方法框架/流程图": [
        "已完成publication-ready流程图，并将按以下主线组织：官方研究区和灾害库存 → 2018因子数据库 → VIF/相关性筛选 → Spatial-CV模型族 → Spatial-CV Stacked Ensemble → SHAP/Calibration/AOA → 域间迁移与机制对比 → 2017-2024动态因子时序压力测试 → CPEC/KKH基础设施暴露和强降雨/强震情景。",
        "后续会将该流程图作为方法部分核心图，并把模型评价、可靠性、解释性和应用输出放在同一逻辑框架中。"
    ],
    "4. 使用官方数据": [
        "已落实。研究区不再使用早期KKH小范围子集，而采用官方CPEC研究区边界：Pakistan + Xinjiang Uygur, Kashgar。GEE资产为 projects/ee-mohibullah141300/assets/cpec_boundary_official_study_area，本地边界也已保存到项目中。",
        "子区域/子域划分将基于官方行政边界和走廊地貌差异共同定义，当前公开名称为 Kashgar (Xinjiang, China)、Gilgit-Baltistan、KP-AJK、Balochistan、Punjab-Sindh lowland corridor；FATA已按现行行政关系并入KP，不再单独作为公开子域。"
    ],
    "5. 进一步整理出适用于动态易发性评价的数据集": [
        "修订后的时序策略为：不把论文做成2017-2024八个独立年度模型的地图集，而是在严格验证的2018模型基础上，开展2017-2024 temporal transfer / dynamic susceptibility stress-testing。这样更符合目前可获得的年度标签条件，也更容易向审稿人解释。",
        "2017-2024作为主时窗的原因是AlphaEarth年度嵌入从2017年开始；2013-2016如导师需要，可作为Conventional-only补充敏感性分析，但不作为AlphaEarth主创新结果。",
        "年度变化因子包括monsoon rainfall total、extreme rainfall metric、NDVI median、NDVI amplitude、land cover、AlphaEarth 64 bands；静态/半静态因子如DEM派生因子、岩性、土壤、道路/河流/断层距离保持固定。PGA将作为强震情景或年度地震触发扩展因子加入。"
    ],
    "6.开展研究区易发性的其时空规律分析": [
        "已完成2018 Spatial-CV和leave-one-domain-out域间迁移实验。该部分用于证明模型不是只在随机样本上表现好，而是在不同地貌/行政子域之间也具备可解释的泛化能力。",
        "当前域间结果将解释为：不同子域ROC-AUC差异反映了孕灾环境异质性、样本数量差异和训练域-测试域相似度差异；AOA/transfer confidence用于说明哪些区域属于可靠插值，哪些区域属于外推预测。",
        "后续动态部分将把2017-2024年度预测压缩为mean susceptibility、temporal variability、trend、persistent high-susceptibility zones、largest year-to-year change和KKH/CPEC road exposure，而不是把所有年度地图都放入主文。"
    ],
    "7. 补充实验": [
        "下一步将把强降雨和强震作为情景分支，而不是重新训练主模型。强降雨情景将用极端降雨指标（如历史高分位或百年一遇设计降雨，视数据可得性确定）替换/扰动年度降雨因子；强震情景将用PGA或断层地震情景因子扰动地震触发条件。",
        "输出将包括：scenario probability、scenario minus baseline、reliable high-risk zones、uncertain high-risk zones，以及中巴铁路/道路、油气管线、Gwadar港等关键基础设施的暴露长度/热点段。该部分需要最终确认铁路、管线和港口矢量数据源。"
    ],
    "1. 论文写作过程中": [
        "写作中将严格区分generalization和robustness。Generalization用于描述空间交叉验证或留一子域测试中的外部预测能力；robustness用于描述模型在不同采样、特征集、子域、AOA可靠性和情景扰动下结果是否稳定。",
        "因此原来类似“Small performance gap suggests good generalization”的表述将改为更精确的句子，例如：small domain-transfer performance loss indicates spatial generalization, while stable rankings across feature sets/sampling/scenarios support robustness。"
    ],
    "2. 阅读《Geoscience Frontiers》": [
        "写作风格将向RSE、Geoscience Frontiers和Journal of Earth Science靠拢，但主线会避免普通模型比较论文的写法。当前拟突出四个贡献：AlphaEarth added value、spatial transferability/AOA reliability、TreeSHAP mechanism interpretation、CPEC/KKH infrastructure exposure and scenario relevance。",
        "建议投稿顺序同意暂定为 RSE → Geoscience Frontiers → Journal of Earth Science。初稿完成后再根据创新强度、图件质量和动态/情景实验完整度与导师进一步确定目标期刊。"
    ],
    "参考文献": [
        "除老师给出的Gao et al. (2024)、Jones et al. (2021)、Lee et al. (2022)、Liu et al. (2023)外，方法部分还将补充空间交叉验证、Area of Applicability、TreeSHAP、stacked ensemble、AlphaEarth/Satellite Embedding和最新transfer learning/large-area LSM文献，用于支撑每一步方法选择。",
        "参考文献不会只堆砌在文末，而会对应到具体方法步骤：因子选择、动态因子构建、空间验证、域间迁移、模型解释、可靠性/外推识别和基础设施暴露分析。"
    ],
}


def main() -> None:
    OUT_MAIN.parent.mkdir(parents=True, exist_ok=True)
    OUT_COPY.parent.mkdir(parents=True, exist_ok=True)

    # Work from a byte-for-byte copy first; then python-docx inserts only the
    # response paragraphs, preserving the original document structure/images.
    temp_path = OUT_MAIN.with_suffix(".working.docx")
    shutil.copy2(SOURCE, temp_path)

    doc = Document(str(temp_path))
    for starts_with, lines in RESPONSES.items():
        paragraph = find_paragraph(doc, starts_with)
        add_response_after(paragraph, lines)

    doc.save(str(OUT_MAIN))
    shutil.copy2(OUT_MAIN, OUT_COPY)
    temp_path.unlink(missing_ok=True)
    print(OUT_MAIN)
    print(OUT_COPY)


if __name__ == "__main__":
    main()
