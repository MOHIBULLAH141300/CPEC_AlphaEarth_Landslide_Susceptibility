"""Create a supervisor-ready Word summary of the current 2018 V3 results."""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT = Path(r"D:\DING PROJECT")
PACKAGE = PROJECT / "00_READ_ME_FIRST_2018_V3_RESULTS"
REPORTS = PROJECT / "05_reports"
OUT = PACKAGE / "CPEC_KKH_2018_V3_current_model_results_supervisor_summary_2026-05-11.docx"
REPORT_COPY = REPORTS / OUT.name

PERF = PACKAGE / "02_model_performance_tables" / "clean_model_performance_summary_for_supervisor.csv"
DOMAIN = (
    PACKAGE
    / "10_domain_specific_2018_results"
    / "00_domain_comparison_summary"
    / "domain_transferability_mechanism_comparison_table.csv"
)
GROUP = (
    PACKAGE
    / "10_domain_specific_2018_results"
    / "00_domain_comparison_summary"
    / "domainwise_treeshap_group_importance_all_feature_sets.csv"
)
VIF = PACKAGE / "06_samples_factors_and_vif" / "v3_final_selected_vif.csv"

FLOWCHART = (
    PACKAGE
    / "03_figures"
    / "04_methodology_flowchart"
    / "figure_cpec_kkh_publishable_methodology_flowchart_2026-05-10.png"
)
TRANSFER_HEATMAP = (
    PACKAGE
    / "10_domain_specific_2018_results"
    / "00_domain_comparison_summary"
    / "figure_domain_transferability_reliability_comparison_heatmap.png"
)
TOP_FACTOR_HEATMAP = (
    PACKAGE
    / "10_domain_specific_2018_results"
    / "00_domain_comparison_summary"
    / "figure_domain_top_factor_shap_heatmap_fused.png"
)
GROUP_HEATMAP = (
    PACKAGE
    / "10_domain_specific_2018_results"
    / "00_domain_comparison_summary"
    / "figure_domain_mechanism_group_heatmap_fused.png"
)


ACCENT = RGBColor(31, 78, 121)
MUTED = RGBColor(95, 95, 95)
LIGHT = "D9EAF7"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False, color: RGBColor | None = None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(9)
    if color is not None:
        run.font.color.rgb = color
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "B7B7B7")


def set_table_width(table, widths_in: list[float]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for row in table.rows:
        for cell, width in zip(row.cells, widths_in):
            cell.width = Inches(width)
    set_table_borders(table)


def add_table(doc: Document, headers: list[str], rows: list[list[object]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for cell, header in zip(hdr.cells, headers):
        set_cell_shading(cell, LIGHT)
        set_cell_text(cell, header, bold=True, color=ACCENT)
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            set_cell_text(cell, value)
    set_table_width(table, widths)
    doc.add_paragraph()


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(5)
    p.add_run(text)


def add_numbered(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(5)
    p.add_run(text)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED


def add_picture_if_exists(doc: Document, path: Path, width_in: float, caption: str) -> None:
    if not path.exists():
        p = doc.add_paragraph()
        run = p.add_run(f"[Figure missing: {path}]")
        run.italic = True
        run.font.color.rgb = RGBColor(160, 0, 0)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width_in))
    add_caption(doc, caption)


def fmt(x: float | int | str | None, decimals: int = 3) -> str:
    if x is None or pd.isna(x):
        return "-"
    if isinstance(x, str):
        return x
    return f"{float(x):.{decimals}f}"


def configure_document(doc: Document) -> None:
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.left_margin = Inches(0.85)
    sec.right_margin = Inches(0.85)
    sec.top_margin = Inches(0.8)
    sec.bottom_margin = Inches(0.8)
    sec.header_distance = Inches(0.35)
    sec.footer_distance = Inches(0.35)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.line_spacing = 1.08
    styles["Normal"].paragraph_format.space_after = Pt(5)
    styles["Title"].font.name = "Arial"
    styles["Title"].font.size = Pt(22)
    styles["Title"].font.bold = True
    styles["Title"].font.color.rgb = ACCENT
    for name, size in [("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11)]:
        styles[name].font.name = "Arial"
        styles[name].font.size = Pt(size)
        styles[name].font.bold = True
        styles[name].font.color.rgb = ACCENT
        styles[name].paragraph_format.space_before = Pt(10)
        styles[name].paragraph_format.space_after = Pt(4)

    header = sec.header.paragraphs[0]
    header.text = "CPEC/KKH landslide susceptibility - 2018 V3 supervisor summary"
    header.runs[0].font.name = "Arial"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = MUTED

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.text = "Prepared from D:\\DING PROJECT clean 2018 V3 results"
    footer.runs[0].font.name = "Arial"
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = MUTED


def add_title_block(doc: Document) -> None:
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.add_run("CPEC/KKH Landslide Susceptibility Modelling")
    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(10)
    r = sub.add_run("Current 2018 V3 model structure, results, domain transferability, and next steps")
    r.font.size = Pt(12)
    r.font.color.rgb = MUTED

    add_table(
        doc,
        ["Item", "Summary"],
        [
            ["Study focus", "Landslide susceptibility for the CPEC region with special attention to KKH/CPEC infrastructure."],
            ["Main baseline year", "2018, because AlphaEarth annual embeddings are available from 2017 onward and 2018 is the validated baseline comparison year."],
            ["Main model", "Spatial-CV Stacked Ensemble evaluated against Logistic Regression, Random Forest, Extra Trees, XGBoost, LightGBM, and CatBoost."],
            ["Main feature sets", "Conventional; AlphaEarth Embeddings; Conventional + AlphaEarth Embeddings."],
            ["Current status", "Core 2018 workflow is complete. Temporal 2017-2024 extension remains planned and partly staged."],
            ["Prepared", date.today().isoformat()],
        ],
        [1.6, 5.7],
    )


def make_doc() -> Document:
    perf = pd.read_csv(PERF)
    domain = pd.read_csv(DOMAIN)
    group = pd.read_csv(GROUP)
    vif = pd.read_csv(VIF)

    doc = Document()
    configure_document(doc)
    add_title_block(doc)

    doc.add_heading("1. Executive Summary", level=1)
    for text in [
        "The project has produced a complete 2018 V3 landslide susceptibility framework for the official CPEC study area.",
        "The strongest result is not only high accuracy; it is a transferability-aware, explainable, reliability-aware model structure.",
        "The fused feature set, Conventional + AlphaEarth Embeddings, gives the best overall balance of discrimination, calibration, and spatial transferability.",
        "The domain analysis shows where the model generalizes well, where uncertainty is higher, and which controlling mechanisms dominate each subdomain.",
    ]:
        add_bullet(doc, text)

    doc.add_heading("2. Current Model Structure", level=1)
    doc.add_paragraph(
        "The modelling design uses one validated CPEC-wide framework rather than separate regional models. "
        "This is important because it allows fair comparison among subdomains and directly tests spatial generalization."
    )
    add_table(
        doc,
        ["Component", "What was done", "Why it matters"],
        [
            ["Official boundary", "All samples and maps are clipped to the official CPEC study boundary.", "Keeps the analysis spatially consistent."],
            ["Corrected factor database", "Terrain, hydro-climate, vegetation, land cover, geology, soil, seismicity, road distance, river distance, and fault distance were included.", "Covers physical, environmental, and infrastructure controls."],
            ["VIF gate", f"All final conventional factors passed global VIF screening; maximum final VIF = {vif['vif'].max():.2f}.", "Reduces unstable or redundant predictors before modelling."],
            ["Model family", "Seven model types were evaluated with spatial CV.", "Shows that the chosen model is not an arbitrary single-model result."],
            ["Stacking", "Base learners are combined through a Spatial-CV Stacked Ensemble.", "Uses complementary model strengths while reducing reliance on any one learner."],
            ["Domain tests", "Leave-one-domain-out tests were run for five CPEC subdomains.", "Tests spatial transferability to withheld regions."],
            ["AOA", "Area of Applicability and transfer confidence were generated.", "Separates reliable predictions from extrapolated predictions."],
            ["Explainability", "SHAP and TreeSHAP summaries were created globally and by domain.", "Provides process interpretation, not just accuracy numbers."],
        ],
        [1.35, 3.0, 2.95],
    )

    doc.add_heading("3. Model Performance", level=1)
    doc.add_paragraph(
        "The fused feature set produced the best overall performance. XGBoost and CatBoost gave the highest raw AUC values, "
        "while the Spatial-CV Stacked Ensemble gave the strongest balanced classification and lowest Brier score among fused models."
    )
    selected = perf[
        ((perf["Feature set"] == "Conventional + AlphaEarth Embeddings") & (perf["Model"].isin(["Spatial-CV Stacked Ensemble", "XGBoost", "CatBoost"])))
        | ((perf["Feature set"].isin(["Conventional", "AlphaEarth Embeddings"])) & (perf["Model"] == "Spatial-CV Stacked Ensemble"))
    ].copy()
    rows = []
    for _, r in selected.iterrows():
        rows.append(
            [
                r["Feature set"],
                r["Model"],
                fmt(r["ROC-AUC mean"]),
                fmt(r["PR-AUC mean"]),
                fmt(r["Balanced accuracy mean"]),
                fmt(r["F1 mean"]),
                fmt(r["Brier score mean"]),
            ]
        )
    add_table(
        doc,
        ["Feature set", "Model", "ROC-AUC", "PR-AUC", "Bal. acc.", "F1", "Brier"],
        rows,
        [1.9, 1.55, 0.7, 0.7, 0.75, 0.65, 0.65],
    )
    doc.add_paragraph(
        "Interpretation: the fused model improves over the conventional-only and embedding-only structures, showing that AlphaEarth provides complementary information rather than duplicating conventional factors."
    )

    doc.add_page_break()
    doc.add_heading("4. Domain Transferability", level=1)
    doc.add_paragraph(
        "Leave-one-domain-out means the model is trained on all other CPEC domains and tested only on the withheld domain. "
        "The values below therefore measure spatial generalization rather than ordinary random-split accuracy."
    )
    drows = []
    for _, r in domain.iterrows():
        drows.append(
            [
                r["domain"],
                fmt(r["roc_auc_conventional_and_alphaearth_embeddings"]),
                fmt(r["pr_auc_ap_conventional_and_alphaearth_embeddings"]),
                fmt(r["f1_conventional_and_alphaearth_embeddings"]),
                fmt(r["brier_conventional_and_alphaearth_embeddings"]),
                f"{r['fused_aoa_coverage_percent']:.1f}%",
                f"{r['uncertain_high_area_percent']:.1f}%",
            ]
        )
    add_table(
        doc,
        ["Held-out domain", "Fused AUC", "Fused PR-AUC", "F1", "Brier", "AOA", "Uncertain high"],
        drows,
        [2.2, 0.75, 0.8, 0.6, 0.65, 0.65, 0.9],
    )
    doc.add_paragraph(
        "The fused model transferred strongly to all domains. Balochistan has the lowest AUC, which indicates that it is the most difficult transfer domain and likely has more distinct geomorphic or environmental controls. "
        "Punjab-Sindh has the highest AUC, but this should be interpreted with its simpler lowland contrast and smaller landslide-prone footprint in mind."
    )

    add_picture_if_exists(
        doc,
        TRANSFER_HEATMAP,
        6.8,
        "Figure 1. Domain transferability, Area of Applicability, and uncertainty comparison.",
    )

    doc.add_page_break()
    doc.add_heading("5. Area of Applicability and Reliability", level=1)
    doc.add_paragraph(
        "ROC-AUC and PR-AUC evaluate ranking skill at sample points. Area of Applicability evaluates whether each mapped pixel is similar enough to the training data for the prediction to be trusted. "
        "This is why AOA is essential for a large corridor such as CPEC: a model can predict everywhere, but some locations are extrapolations."
    )
    for text in [
        "High probability inside AOA: stronger candidate for susceptibility interpretation and planning.",
        "High probability outside AOA: possible hotspot, but should be treated as uncertain and prioritized for field verification.",
        "Reliability-weighted probability is probability multiplied by transfer confidence; it is a secondary confidence-adjusted product, not the raw susceptibility probability map.",
        "Conventional reliability-weighted maps can look near zero where transfer confidence is low; this is expected and documented separately.",
    ]:
        add_bullet(doc, text)

    doc.add_heading("6. Domain Mechanisms From TreeSHAP", level=1)
    doc.add_paragraph(
        "The domain-wise TreeSHAP analysis explains the XGBoost base learner within each feature set. "
        "It is used to identify dominant conditioning mechanisms while the Spatial-CV Stacked Ensemble remains the main predictive model."
    )
    fused_groups = group[group["feature_set"] == "Conventional + AlphaEarth Embeddings"].copy()
    mech_rows = []
    for domain_name, d in fused_groups.groupby("domain", sort=False):
        top = d.sort_values("group_relative_importance_percent", ascending=False).head(2)
        mech_rows.append(
            [
                domain_name,
                "; ".join(
                    f"{rr['feature_group']} ({rr['group_relative_importance_percent']:.1f}%)"
                    for _, rr in top.iterrows()
                ),
            ]
        )
    add_table(doc, ["Domain", "Dominant fused-model mechanism groups"], mech_rows, [2.3, 5.0])
    add_picture_if_exists(
        doc,
        TOP_FACTOR_HEATMAP,
        6.8,
        "Figure 2. Domain-specific top TreeSHAP factors for the fused model.",
    )
    add_picture_if_exists(
        doc,
        GROUP_HEATMAP,
        6.8,
        "Figure 3. Domain mechanism groups from fused-model TreeSHAP importance.",
    )

    doc.add_heading("7. Scientific Meaning of the Current Results", level=1)
    for text in [
        "The framework shows that AlphaEarth embeddings add useful surface-condition information to conventional landslide factors.",
        "The fused model is not only accurate; it is spatially tested through leave-one-domain-out validation.",
        "AOA and transfer-confidence maps show where predictions are reliable versus extrapolated.",
        "Domain-wise SHAP shows that the same CPEC-wide model can have different controlling mechanisms in different regions.",
        "Road exposure outputs turn susceptibility modelling into an infrastructure-relevant CPEC/KKH product.",
    ]:
        add_numbered(doc, text)

    doc.add_heading("8. Main Innovation", level=1)
    p = doc.add_paragraph()
    p.add_run("Proposed contribution: ").bold = True
    p.add_run(
        "a transferability-aware, explainable, AlphaEarth-enhanced landslide susceptibility framework for the CPEC/KKH corridor."
    )
    doc.add_paragraph(
        "This is stronger than a standard model-comparison paper because the workflow evaluates accuracy, spatial generalization, model applicability, factor mechanisms, reliability, and infrastructure exposure together."
    )

    doc.add_page_break()
    doc.add_heading("9. Remaining Work", level=1)
    add_table(
        doc,
        ["Task", "Status", "Recommended action"],
        [
            ["2018 core results", "Mostly complete", "Use this as the main paper foundation."],
            ["Figure selection", "Needs final polish", "Select the few strongest figures for main text and move the rest to supplementary material."],
            ["Temporal extension", "Planned and partly staged", "Use fixed 2018 model for compact 2017-2024 temporal stress-testing, not eight separate papers."],
            ["2020-2024 annual inputs", "Pending", "Export/download yearly GEE predictor stacks one year at a time."],
            ["Manuscript writing", "Not yet complete", "Convert this summary into Methods, Results, and Discussion sections."],
        ],
        [2.0, 1.4, 3.9],
    )

    doc.add_heading("10. Files to Open First", level=1)
    for text in [
        "README_START_HERE.md in the clean 2018 V3 results package.",
        "Domain comparison summary: domainwise_mechanism_and_transferability_interpretation_report.md.",
        "Temporal strategy: better_temporal_evaluation_strategy_2026-05-10.md.",
    ]:
        add_bullet(doc, text)

    doc.add_heading("11. Short Reference Basis", level=1)
    for text in [
        "Spatial/grouped CV reduces overoptimistic random-split accuracy.",
        "AOA links prediction reliability to similarity with the training space.",
        "TreeSHAP supports global and local interpretation for tree learners.",
        "AlphaEarth annual embeddings justify the 2017-onward temporal extension.",
    ]:
        add_bullet(doc, text)

    return doc


def main() -> None:
    doc = make_doc()
    OUT.parent.mkdir(parents=True, exist_ok=True)
    REPORTS.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    doc.save(REPORT_COPY)
    print(OUT)
    print(REPORT_COPY)


if __name__ == "__main__":
    main()
