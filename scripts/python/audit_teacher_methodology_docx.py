from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


DOCX = Path(
    r"D:\DING PROJECT\00_READ_ME_FIRST_2018_V3_RESULTS"
    r"\CPEC_KKH_updated_next_methodology_summary_for_teacher_2026-05-10.docx"
)
OUT = Path(
    r"D:\DING PROJECT\05_reports\docx_render_qa"
    r"\teacher_methodology_summary_2026-05-10\structural_docx_audit.json"
)


def dxa_value(element, xpath: str) -> int | None:
    found = element.find(xpath)
    if found is None:
        return None
    val = found.get(qn("w:w"))
    return int(val) if val and val.isdigit() else None


def table_width(table) -> dict[str, object]:
    tbl = table._tbl
    tbl_w = tbl.tblPr.find(qn("w:tblW"))
    declared = int(tbl_w.get(qn("w:w"))) if tbl_w is not None and tbl_w.get(qn("w:w"), "").isdigit() else None
    grid = tbl.tblGrid
    grid_widths = []
    if grid is not None:
        for col in grid.findall(qn("w:gridCol")):
            val = col.get(qn("w:w"))
            grid_widths.append(int(val) if val and val.isdigit() else None)
    return {
        "rows": len(table.rows),
        "cols": len(table.columns),
        "declared_width_dxa": declared,
        "grid_widths_dxa": grid_widths,
        "grid_sum_dxa": sum(v for v in grid_widths if isinstance(v, int)),
        "declared_matches_grid": declared == sum(v for v in grid_widths if isinstance(v, int)),
    }


def main() -> None:
    doc = Document(DOCX)
    section = doc.sections[0]
    page_width = section.page_width.twips
    content_width = page_width - section.left_margin.twips - section.right_margin.twips
    paragraphs = list(doc.paragraphs)
    list_paragraphs = [
        p.text for p in paragraphs if p.style and p.style.name in {"List Bullet", "List Number"}
    ]
    fake_list_candidates = [
        p.text for p in paragraphs if p.text.strip().startswith(("-", "*", "1)"))
    ]
    tables = [table_width(t) for t in doc.tables]
    too_wide = [t for t in tables if t["grid_sum_dxa"] > content_width]
    mismatched = [t for t in tables if not t["declared_matches_grid"]]
    result = {
        "docx": str(DOCX),
        "paragraph_count": len(paragraphs),
        "table_count": len(doc.tables),
        "section_count": len(doc.sections),
        "page_width_dxa": page_width,
        "left_margin_dxa": section.left_margin.twips,
        "right_margin_dxa": section.right_margin.twips,
        "content_width_dxa": content_width,
        "list_paragraph_count": len(list_paragraphs),
        "fake_list_candidate_count": len(fake_list_candidates),
        "fake_list_candidates": fake_list_candidates[:10],
        "table_audit": tables,
        "too_wide_table_count": len(too_wide),
        "mismatched_table_width_count": len(mismatched),
        "passed_structural_audit": len(too_wide) == 0 and len(mismatched) == 0 and len(fake_list_candidates) == 0,
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(result, indent=2), encoding="utf-8")
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
