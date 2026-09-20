"""Build a DOCX for the CPEC methods/discussion/reviewer package.

The source is a manuscript-ready Markdown file. The builder applies a compact
formal Word style based on the Documents skill standard_business_brief preset.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(r"D:\DING PROJECT")
MD_PATH = (
    PROJECT_ROOT
    / "05_reports"
    / "manuscript_drafts"
    / "geoscience_frontiers_methods_discussion_reviewer_package_2026-06-07.md"
)
OUT_DOCX = (
    PROJECT_ROOT
    / "05_reports"
    / "manuscript_drafts"
    / "geoscience_frontiers_methods_discussion_reviewer_package_2026-06-07.docx"
)

CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def set_cell_margins(table, top=80, start=120, bottom=80, end=120):
    tbl_pr = table._tbl.tblPr
    tbl_cell_mar = tbl_pr.first_child_found_in("w:tblCellMar")
    if tbl_cell_mar is None:
        tbl_cell_mar = OxmlElement("w:tblCellMar")
        tbl_pr.append(tbl_cell_mar)
    for margin, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tbl_cell_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tbl_cell_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9DEE7"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def shade_cell(cell, fill="F2F4F7"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def keep_table_row_together(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)


def set_table_fixed_width(table, widths_dxa: list[int]):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.10):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def clean_inline(text: str) -> str:
    text = text.replace("`", "")
    text = text.replace("**", "")
    text = text.replace("\\(", "(").replace("\\)", ")")
    text = text.replace("\\[", "[").replace("\\]", "]")
    text = text.replace("\\Delta", "Delta")
    text = text.replace("\\times", "x")
    text = text.replace("\\ge", ">=")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def split_table_row(line: str) -> list[str]:
    stripped = line.strip().strip("|")
    return [clean_inline(cell.strip()) for cell in stripped.split("|")]


def is_separator_row(line: str) -> bool:
    cells = split_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell.strip()) for cell in cells)


def column_widths(ncols: int) -> list[int]:
    if ncols == 3:
        return [700, 3200, 5460]
    if ncols == 4:
        return [700, 2200, 3950, 2510]
    if ncols == 5:
        return [520, 2350, 2700, 2600, 1190]
    if ncols == 6:
        return [420, 1700, 2300, 1700, 1850, 1390]
    if ncols == 7:
        return [420, 1300, 1700, 1500, 1500, 1500, 1440]
    return [CONTENT_WIDTH_DXA // ncols] * ncols


def add_table(doc: Document, lines: list[str]):
    rows = [split_table_row(line) for line in lines if not is_separator_row(line)]
    if not rows:
        return
    ncols = max(len(row) for row in rows)
    for row in rows:
        row.extend([""] * (ncols - len(row)))

    table = doc.add_table(rows=len(rows), cols=ncols)
    set_table_fixed_width(table, column_widths(ncols))
    set_cell_margins(table)
    set_table_borders(table)
    table.style = "Table Grid"
    for r_idx, row in enumerate(rows):
        keep_table_row_together(table.rows[r_idx])
        if r_idx == 0:
            repeat_table_header(table.rows[r_idx])
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx == 0 or r_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(value)
            run.font.name = "Calibri"
            run.font.size = Pt(8.4 if ncols >= 5 else 9)
            if r_idx == 0:
                run.bold = True
                shade_cell(cell)
            set_paragraph_spacing(p, before=0, after=0, line=1.0)
    doc.add_paragraph()


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.text = "Methods, Discussion, and Reviewer-Defense Package"
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in p.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(102, 102, 102)


def setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    add_footer(section)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.10

    return doc


def build_docx():
    doc = setup_document()
    lines = MD_PATH.read_text(encoding="utf-8").splitlines()
    para_lines: list[str] = []
    table_lines: list[str] = []
    first_heading = True

    def flush_para():
        nonlocal para_lines
        if not para_lines:
            return
        text = clean_inline(" ".join(para_lines))
        para_lines = []
        if not text:
            return
        if text in {"[", "]"}:
            return
        p = doc.add_paragraph()
        if text.startswith("Delta ") or text.startswith("R =") or text.startswith("where "):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if not text.startswith("where ") else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(text)
            run.font.name = "Cambria Math"
            run.font.size = Pt(10.5)
        else:
            run = p.add_run(text)
            run.font.name = "Calibri"
            run.font.size = Pt(11)
        set_paragraph_spacing(p, before=0, after=6, line=1.10)

    def flush_table():
        nonlocal table_lines
        if table_lines:
            add_table(doc, table_lines)
            table_lines = []

    for line in lines:
        stripped = line.strip()
        if stripped.startswith("|"):
            flush_para()
            table_lines.append(stripped)
            continue
        flush_table()

        if not stripped:
            flush_para()
            continue

        if stripped.startswith("# "):
            flush_para()
            text = clean_inline(stripped[2:])
            if first_heading:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run(text)
                run.bold = True
                run.font.name = "Calibri"
                run.font.size = Pt(18)
                run.font.color.rgb = RGBColor.from_string("0B2545")
                set_paragraph_spacing(p, before=0, after=10, line=1.10)
                first_heading = False
            else:
                doc.add_heading(text, level=1)
            continue
        if stripped.startswith("## "):
            flush_para()
            doc.add_heading(clean_inline(stripped[3:]), level=1)
            continue
        if stripped.startswith("### "):
            flush_para()
            doc.add_heading(clean_inline(stripped[4:]), level=2)
            continue
        if stripped.startswith("- "):
            flush_para()
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(clean_inline(stripped[2:]))
            set_paragraph_spacing(p, before=0, after=4, line=1.10)
            continue
        para_lines.append(stripped)

    flush_para()
    flush_table()

    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_DOCX)
    print(OUT_DOCX)


if __name__ == "__main__":
    build_docx()
