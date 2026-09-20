"""Build highlights and a reproducible graphical abstract for manuscript v3."""

from __future__ import annotations

from pathlib import Path

import matplotlib as mpl
import matplotlib.pyplot as plt
import numpy as np
import rasterio
import shapefile
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from matplotlib.collections import LineCollection
from matplotlib.colors import PowerNorm
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch


PROJECT = Path(r"D:\DING PROJECT")
PAPER = PROJECT / "FINAL PAPER" / "Manuscript_single_file" / "v3_2026-07-17"
RASTER = (
    PROJECT
    / "04_maps"
    / "manuscript_v3_baseline_susceptibility_scores_250m"
    / "cpec_baseline_conventional_alphaearth_embeddings_stacked_susceptibility_score_250m.tif"
)
LAYER_DIR = PROJECT / "FINAL PAPER" / "00_Study_Area" / "Figure_1_ArcMap_layers"
BOUNDARY = LAYER_DIR / "01_CPEC_study_boundary.shp"
KKH = LAYER_DIR / "04_KKH_route.shp"

HIGHLIGHTS = [
    "Nested buffered Spatial-CV prevents leakage in corridor-scale susceptibility.",
    "AlphaEarth fusion improves ranking and Brier score over conventional factors.",
    "Matched controls and typology tests quantify inventory and sampling effects.",
    "LODO and harmonised AoA expose non-uniform transfer across CPEC.",
    "Road-distance ablation preserves KKH hotspot rankings (rho = 0.981).",
]


def shape_parts(shape):
    points = np.asarray(shape.points, dtype=float)
    starts = list(shape.parts) + [len(points)]
    for start, end in zip(starts[:-1], starts[1:]):
        if end > start:
            yield points[start:end]


def line_segments(path: Path):
    reader = shapefile.Reader(str(path), encoding="latin1")
    segments = []
    for shape in reader.iterShapes():
        segments.extend([part for part in shape_parts(shape) if len(part) >= 2])
    return segments


def build_highlights():
    for item in HIGHLIGHTS:
        if len(item) > 85:
            raise ValueError(f"Highlight exceeds 85 characters ({len(item)}): {item}")

    txt = PAPER / "CPEC_manuscript_v3_highlights_2026-07-17.txt"
    txt.write_text("\n".join(f"- {item}" for item in HIGHLIGHTS) + "\n", encoding="utf-8")

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(11)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Highlights")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(16)
    for item in HIGHLIGHTS:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(7)
        p.add_run(item)
    output = PAPER / "CPEC_manuscript_v3_highlights_2026-07-17.docx"
    doc.save(output)


def module(ax, x, y, w, h, title, lines, color):
    patch = FancyBboxPatch(
        (x, y),
        w,
        h,
        boxstyle="round,pad=0.012,rounding_size=0.012",
        facecolor=color,
        edgecolor="#303030",
        linewidth=1.1,
        transform=ax.transAxes,
        zorder=3,
    )
    ax.add_patch(patch)
    ax.text(x + 0.03 * w, y + 0.82 * h, title, transform=ax.transAxes, fontsize=11.5, fontweight="bold", va="center", zorder=4)
    ax.text(x + 0.05 * w, y + 0.44 * h, lines, transform=ax.transAxes, fontsize=9.4, va="center", linespacing=1.45, zorder=4)


def build_graphical_abstract():
    mpl.rcParams.update({"font.family": "Arial", "font.size": 10, "pdf.fonttype": 42})
    fig = plt.figure(figsize=(12.8, 4.8), facecolor="white")
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_axis_off()

    ax.text(
        0.5,
        0.95,
        "Transferability-aware landslide susceptibility for a transboundary corridor",
        ha="center",
        va="center",
        fontsize=16,
        fontweight="bold",
        color="#1B1B1B",
    )

    module(
        ax,
        0.025,
        0.19,
        0.22,
        0.64,
        "1  Data and representations",
        "1,508 mapped slope failures\n1,808 controls\n\n19 conventional factors\n64 AlphaEarth embeddings",
        "#DDEBF7",
    )
    module(
        ax,
        0.285,
        0.19,
        0.29,
        0.64,
        "2  Leakage-resistant evaluation",
        "Nested 5-fold Spatial-CV\n20 km outer and inner buffers\n6 base learners + L2 stack\n\nMatched controls and typology tests\nLODO transfer + harmonised AoA",
        "#E3F1E8",
    )

    for x1, x2 in ((0.245, 0.285), (0.575, 0.615)):
        ax.add_patch(
            FancyArrowPatch(
                (x1, 0.51),
                (x2, 0.51),
                transform=ax.transAxes,
                arrowstyle="-|>",
                mutation_scale=16,
                linewidth=1.7,
                color="#3A3A3A",
            )
        )

    # Output module and actual fused score map.
    output_box = FancyBboxPatch(
        (0.615, 0.19),
        0.36,
        0.64,
        boxstyle="round,pad=0.012,rounding_size=0.012",
        facecolor="#F8F1D8",
        edgecolor="#303030",
        linewidth=1.1,
        transform=ax.transAxes,
        zorder=2,
    )
    ax.add_patch(output_box)
    ax.text(0.63, 0.77, "3  Transfer and road-prioritisation evidence", transform=ax.transAxes, fontsize=11.5, fontweight="bold", zorder=4)

    map_ax = fig.add_axes([0.625, 0.255, 0.16, 0.46])
    with rasterio.open(RASTER) as src:
        factor = max(1, int(max(src.width, src.height) / 1300))
        data = src.read(1, out_shape=(max(1, src.height // factor), max(1, src.width // factor)), masked=True)
        extent = (src.bounds.left, src.bounds.right, src.bounds.bottom, src.bounds.top)
    map_ax.imshow(data, extent=extent, origin="upper", cmap="magma", norm=PowerNorm(gamma=0.55, vmin=0, vmax=1), interpolation="bilinear")
    map_ax.add_collection(LineCollection(line_segments(BOUNDARY), colors="#111111", linewidths=0.8, zorder=3))
    map_ax.add_collection(LineCollection(line_segments(KKH), colors="#00FFFF", linewidths=1.4, zorder=4))
    map_ax.set_xlim(59.8, 80.6)
    map_ax.set_ylim(22.9, 42.0)
    map_ax.set_aspect("equal")
    map_ax.set_xticks([])
    map_ax.set_yticks([])
    for spine in map_ax.spines.values():
        spine.set_linewidth(0.7)
    map_ax.text(0.04, 0.04, "Fused score + KKH", transform=map_ax.transAxes, fontsize=8.4, color="white", fontweight="bold", path_effects=[])

    # Keep the narrow evidence column as four clearly separated reading blocks.
    ax.text(0.805, 0.690, "Nested Spatial-CV", transform=ax.transAxes, fontsize=8.9, fontweight="bold", va="top")
    ax.text(0.805, 0.625, "ROC-AUC  0.967\nPR-AUC    0.960\nBrier       0.065", transform=ax.transAxes, fontsize=8.9, linespacing=1.45, va="top")
    ax.text(0.805, 0.455, "Held-out domains", transform=ax.transAxes, fontsize=8.9, fontweight="bold", va="top")
    ax.text(0.805, 0.395, "ROC-AUC  0.894-0.990\nAoA          40.2%-95.4%", transform=ax.transAxes, fontsize=8.9, linespacing=1.45, va="top")
    ax.text(0.805, 0.295, "KKH above P90: 1,088.4 km", transform=ax.transAxes, fontsize=8.6, fontweight="bold", color="#8C2D04", va="top")
    ax.text(0.805, 0.238, "No-road rank rho = 0.981", transform=ax.transAxes, fontsize=8.6, fontweight="bold", color="#006D5B", va="top")

    ax.text(
        0.5,
        0.08,
        "Foundation-model embeddings are useful when their added value survives spatial leakage controls, sampling tests, domain transfer, and AoA diagnostics.",
        ha="center",
        va="center",
        fontsize=10.5,
        color="#2A2A2A",
    )

    png = PAPER / "CPEC_manuscript_v3_graphical_abstract_2026-07-17.png"
    pdf = PAPER / "CPEC_manuscript_v3_graphical_abstract_2026-07-17.pdf"
    fig.savefig(png, dpi=300, bbox_inches="tight", pad_inches=0.05)
    fig.savefig(pdf, bbox_inches="tight", pad_inches=0.05)
    plt.close(fig)


if __name__ == "__main__":
    build_highlights()
    build_graphical_abstract()
    print(PAPER)
