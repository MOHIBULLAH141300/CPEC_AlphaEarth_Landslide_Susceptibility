"""Run a reproducible structural and terminology audit of manuscript v3.2."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document


STALE_TERMS = {
    "old ROC-AUC": "0.9689",
    "old PR-AUC": "0.9612",
    "old Brier score": "0.0610",
    "old AoA 34.4%": "34.4%",
    "old AoA 95.8%": "95.8%",
    "old AoA 97.7%": "97.7%",
    "old road length 1,078": "1,078",
    "reliability-weighted": "reliability-weighted",
    "Taylor diagram": "Taylor",
    "internal v3 label": "manuscript-v3",
    "excluded temporal period": "2017-2024",
    "excluded temporal period (en dash)": "2017–2024",
    "excluded dated validation": "dated-event temporal validation",
}

PLACEHOLDER_PATTERNS = [
    r"AUTHOR INPUT REQUIRED",
    r"\[Supervisor name\]",
    r"\[Department, institution, city, country\]",
    r"\[email address\]",
    r"\[Co-author name\]",
]


def words(text: str) -> list[str]:
    return re.findall(r"[A-Za-z0-9]+(?:[-–][A-Za-z0-9]+)*", text)


def extract_doc_text(doc: Document) -> str:
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def paragraph_between(doc: Document, start: str, end_prefix: str) -> str:
    collecting = False
    parts: list[str] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text == start:
            collecting = True
            continue
        if collecting and text.startswith(end_prefix):
            break
        if collecting and text:
            parts.append(text)
    return " ".join(parts)


def numbered_items(text: str, label: str) -> list[int]:
    return sorted({int(value) for value in re.findall(rf"\b{label}\s+(\d+)\b", text)})


def run_audit(docx: Path, report: Path) -> str:
    doc = Document(docx)
    text = extract_doc_text(doc)
    abstract = paragraph_between(doc, "Abstract", "Keywords:")
    headings = [
        p.text.strip()
        for p in doc.paragraphs
        if p.text.strip() and p.style and p.style.name.startswith("Heading")
    ]
    table_captions = [p.text.strip() for p in doc.paragraphs if re.match(r"^Table\s+\d+\.", p.text.strip())]
    figure_captions = [p.text.strip() for p in doc.paragraphs if re.match(r"^Figure\s+\d+\.", p.text.strip())]

    stale = {name: token for name, token in STALE_TERMS.items() if token.lower() in text.lower()}
    placeholders = {
        pattern: len(re.findall(pattern, text, flags=re.IGNORECASE))
        for pattern in PLACEHOLDER_PATTERNS
        if re.search(pattern, text, flags=re.IGNORECASE)
    }
    result_text = paragraph_between(doc, "3. Results", "4.")

    expected_tables = list(range(1, 8))
    expected_figures = list(range(1, 10))
    table_numbers = numbered_items("\n".join(table_captions), "Table")
    figure_numbers = numbered_items("\n".join(figure_captions), "Figure")
    checks = {
        "abstract_at_most_250_words": len(words(abstract)) <= 250,
        "seven_tables_numbered_1_to_7": table_numbers == expected_tables,
        "nine_figures_numbered_1_to_9": figure_numbers == expected_figures,
        "seven_results_subsections": sum(h.startswith("3.") and h != "3. Results" for h in headings) == 7,
        "expanded_results_at_least_2000_words": len(words(result_text)) >= 2000,
        "no_stale_analysis_terms": not stale,
        "no_internal_v3_wording": "manuscript-v3" not in text.lower(),
        "no_excluded_temporal_section": "dated-event temporal validation" not in text.lower(),
    }

    lines = [
        "# Manuscript v3.2 final consistency audit",
        "",
        f"- File: `{docx}`",
        f"- Body word count (including references and table text): {len(words(text)):,}",
        f"- Abstract word count: {len(words(abstract))}",
        f"- Results word count: {len(words(result_text)):,}",
        f"- Sections: {len(doc.sections)}",
        f"- Tables: {len(doc.tables)}; numbered captions: {table_numbers}",
        f"- Inline figures: {len(doc.inline_shapes)}; numbered captions: {figure_numbers}",
        "",
        "## Automated checks",
        "",
    ]
    lines.extend(f"- {'PASS' if passed else 'FAIL'}: {name}" for name, passed in checks.items())
    lines.extend(["", "## Remaining author-controlled placeholders", ""])
    if placeholders:
        lines.extend(f"- `{pattern}`: {count}" for pattern, count in placeholders.items())
    else:
        lines.append("- None")
    lines.extend(["", "## Stale-term findings", ""])
    if stale:
        lines.extend(f"- {name}: `{token}`" for name, token in stale.items())
    else:
        lines.append("- None")
    lines.extend(["", "## Heading sequence", ""])
    lines.extend(f"- {heading}" for heading in headings)

    output = "\n".join(lines) + "\n"
    report.parent.mkdir(parents=True, exist_ok=True)
    report.write_text(output, encoding="utf-8")
    return output


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("docx", type=Path)
    parser.add_argument("report", type=Path)
    args = parser.parse_args()
    print(run_audit(args.docx, args.report))


if __name__ == "__main__":
    main()
