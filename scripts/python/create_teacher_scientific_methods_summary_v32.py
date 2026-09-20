"""Create a supervisor brief on the scientific and methodological advances in the current manuscript."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\DING PROJECT\FINAL PAPER\Manuscript_single_file\v3_2_2026-07-17")
OUTPUT = ROOT / "Teacher_brief_scientific_methodological_improvements_v3_2_2026-07-17.docx"

BLUE = "1F4E79"
DARK_BLUE = "17365D"
LIGHT_BLUE = "EAF2F8"
LIGHT_GREEN = "EAF4EF"
MID_GRAY = "666666"
GREEN = "2E7D5A"


def set_font(run, size=10.5, bold=False, color="000000", italic=False):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_font(run, size=9, color=MID_GRAY)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.03
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_font(lead, bold=True, color=DARK_BLUE)
        rest = p.add_run(text[len(bold_lead):])
        set_font(rest)
    else:
        set_font(p.add_run(text))
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.45 + level * 0.22)
    p.paragraph_format.first_line_indent = Inches(-0.22)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.02
    set_font(p.add_run(text))
    return p


def add_callout(doc, title, text, fill=LIGHT_BLUE, accent=BLUE):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.line_spacing = 1.06
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), accent)
    borders.append(left)
    p_pr.append(borders)
    set_font(p.add_run(f"{title}  "), size=10.5, bold=True, color=DARK_BLUE)
    set_font(p.add_run(text), size=10.5, color="263238")
    return p


def add_advance(doc, number, title, improvement, scientific_value, evidence=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    if number == 8:
        p.paragraph_format.page_break_before = True
    set_font(p.add_run(f"{number}. {title}"), size=11.5, bold=True, color=BLUE)
    add_body(doc, f"Methodological improvement: {improvement}", "Methodological improvement:")
    add_body(doc, f"Scientific value: {scientific_value}", "Scientific value:")
    if evidence:
        add_body(doc, f"Evidence now reported: {evidence}", "Evidence now reported:")


def build() -> Path:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.50)
    section.bottom_margin = Inches(0.50)
    section.left_margin = Inches(0.80)
    section.right_margin = Inches(0.80)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.28)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.03
    for name, size, before, after, color in (
        ("Heading 1", 15, 12, 5, BLUE),
        ("Heading 2", 12.5, 8, 4, DARK_BLUE),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    set_font(header.add_run("CPEC landslide susceptibility study | Scientific revision brief"), size=9, bold=True, color=MID_GRAY)
    add_page_number(section.footer.paragraphs[0])

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(3)
    set_font(kicker.add_run("SUPERVISOR REVIEW NOTE"), size=10, bold=True, color=GREEN)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(3)
    set_font(title.add_run("Scientific and methodological strengthening in the current manuscript"), size=21, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(9)
    set_font(
        subtitle.add_run("Transferability-aware landslide susceptibility mapping along the China-Pakistan Economic Corridor"),
        size=12.5,
        color=MID_GRAY,
    )

    metadata = doc.add_paragraph()
    metadata.paragraph_format.space_after = Pt(7)
    for index, (label, value) in enumerate(
        (
            ("Prepared for", "Supervisor review"),
            ("Current manuscript", "v3.2"),
            ("Date", date(2026, 7, 17).strftime("%d %B %Y")),
        )
    ):
        if index:
            set_font(metadata.add_run("   |   "), size=9.3, color=MID_GRAY)
        set_font(metadata.add_run(f"{label}: "), size=9.3, bold=True, color=DARK_BLUE)
        set_font(metadata.add_run(value), size=9.3)

    add_callout(
        doc,
        "Central advance",
        "The manuscript now moves beyond asking which model has the highest AUC. It tests whether susceptibility is spatially transferable, supported by the predictor domain, robust to control sampling and inventory typology, explainable, and useful for prioritising CPEC and Karakoram Highway road segments.",
    )

    add_body(
        doc,
        "Scope of this note: This brief summarises the substantive scientific and methodological strengthening consolidated in the current manuscript relative to the earlier draft shared for review. No numerical result has been altered for this summary.",
        "Scope of this note:",
    )

    add_heading(doc, "1. How the scientific question was strengthened", 1)
    add_body(
        doc,
        "The earlier study concept centred on producing and comparing landslide susceptibility maps. The current manuscript asks a broader and more defensible question: can conventional geoscientific predictors and annual AlphaEarth embeddings be combined to produce susceptibility scores that remain reliable when nearby spatial dependence is removed, when samples and landslide types change, and when the model is transferred across contrasting CPEC subdomains?",
    )
    add_body(
        doc,
        "This reframing changes the contribution from a routine model-ranking exercise into a transferability-aware assessment for a transboundary linear-infrastructure corridor. The mapped values are also defined correctly as case-control susceptibility scores, not annual landslide probabilities or deterministic hazard forecasts.",
    )

    add_heading(doc, "2. Main methodological improvements", 1)

    add_advance(
        doc,
        1,
        "A corrected, physically interpretable predictor foundation",
        "The modelling database was rebuilt across the official CPEC boundary using 19 conventional factors and 64 baseline-year AlphaEarth embedding axes. Terrain derivatives were corrected from the Copernicus DEM, lithology was rerasterised by geological class rather than polygon identifier, and full-CPEC distances to roads, rivers, and active faults were used. Continuous conventional factors passed VIF screening, with the final VIF range restricted to 1.009-3.680.",
        "This removes known data-quality and multicollinearity weaknesses while retaining terrain, hydro-climate, vegetation, land-cover, geological, soil, seismic-background, and proximity controls that have direct geomorphic meaning.",
        "All three 250 m predictor configurations share 19,996,413 valid study-area cells and contain no internal NoData gaps.",
    )

    add_advance(
        doc,
        2,
        "Foundation-model fusion tested as complementarity, not replacement",
        "Three matched feature sets are evaluated: Conventional, AlphaEarth Embeddings, and Conventional + AlphaEarth Embeddings. All 64 embedding axes are retained so that the foundation-model representation is tested without subjective band selection.",
        "The design evaluates whether learned Earth-observation context adds information beyond explicit geoscientific factors. It avoids the unsupported claim that embeddings should replace physical conditioning variables.",
        "Fusion achieved ROC-AUC = 0.967, PR-AUC = 0.960, and Brier score = 0.065. Against Conventional, fusion improved ROC-AUC (p = 0.040) and Brier score (p = 0.015), while the PR-AUC gain was not statistically resolved.",
    )

    add_advance(
        doc,
        3,
        "Leakage-resistant spatial stacking",
        "The official model is a heterogeneous stack of logistic regression, Random Forest, Extra Trees, XGBoost, LightGBM, and CatBoost, combined by an L2-regularised logistic meta-learner. Evaluation uses fully nested five-fold spatial cross-validation. Both inner and outer loops use 1-degree spatial blocks and 20 km exclusion buffers; preprocessing and threshold selection are fitted only within the relevant training data.",
        "This prevents the meta-learner, feature transformations, or nearby spatial neighbours from indirectly accessing the test fold. The reported performance therefore evaluates the complete stacked workflow under geographic separation rather than an optimistic random split.",
    )

    add_advance(
        doc,
        4,
        "Spatially aware uncertainty and calibration",
        "Evaluation was expanded beyond ROC-AUC to PR-AUC, balanced accuracy, F1 score, sensitivity, specificity, precision, Brier score, log loss, and expected calibration error. Confidence intervals and feature-set differences are estimated with 2,000 resamples of 1-degree spatial blocks.",
        "Block resampling respects spatial clustering and permits direct, paired tests of whether apparent improvements are stable across geography. Calibration is interpreted against the sampled case-control labels rather than mislabelled as population occurrence probability.",
    )

    add_advance(
        doc,
        5,
        "Sampling, typology, and road-proximity robustness",
        "Three balanced control sets were matched to positives using elevation, slope, road distance, river distance, fault distance, and monsoon rainfall, with all absolute standardised mean differences below 0.10. Separate landslide-only and rockfall-only models test the pooled slope-failure inventory, and a no-road ablation tests dependence on the strongest accessibility-related factor.",
        "These experiments quantify three major sources of bias that a single high AUC cannot reveal: easy background controls, mixed failure processes, and preferential mapping or disturbance near roads.",
        "Matched repeats retained ROC-AUC = 0.947-0.950. Landslide-only ROC-AUC was 0.911 versus 0.958 for rockfall-only. Removing road distance reduced ROC-AUC to 0.952, but road-segment rankings remained strongly correlated with the full model.",
    )

    add_advance(
        doc,
        6,
        "Direct testing of geographic transferability",
        "Leave-one-domain-out evaluation withholds Kashgar, Gilgit-Baltistan, KPK-AJK, Balochistan, or Punjab-Sindh in turn and trains only on the remaining subdomains, again applying a 20 km exclusion around held-out samples.",
        "This tests whether the framework generalises to an entire unseen geographic domain rather than only to separated points within the pooled study area. It reveals spatial non-stationarity that is hidden by a single global performance value.",
        "Fused ROC-AUC ranged from 0.894 in Balochistan to 0.990 in Punjab-Sindh, demonstrating strong but non-uniform cross-domain transfer.",
    )

    add_advance(
        doc,
        7,
        "Harmonised area-of-applicability diagnostics",
        "Area of applicability (AoA) is calculated for every source-target transfer in the same 15-dimensional whitened predictor space, using a common different-block distance threshold. The same definition is applied to all feature sets.",
        "AoA identifies where target predictor combinations are supported by the source training domain. Separating AoA from accuracy avoids the common mistake of assuming that familiar environmental conditions automatically guarantee good class separation.",
        "Fused AoA ranged from 40.2% in KPK-AJK to 95.4% in Gilgit-Baltistan. AoA coverage was not significantly correlated with ROC-AUC, PR-AUC, Brier score, or calibration error, confirming that applicability and predictive skill are complementary diagnostics.",
    )

    add_advance(
        doc,
        8,
        "Explainability connected to infrastructure decisions",
        "The manuscript separates stack-level meta-learner contributions from exact TreeSHAP explanations of the final XGBoost base learner. Road lines are densified into short segments and evaluated using fused susceptibility, AoA, base-model disagreement, and road-distance ablation.",
        "This preserves physical interpretation while converting the continuous score surface into transparent categories: supported high-score segments and high-score segments requiring field verification. The road overlay is used for prioritisation, not presented as independent validation.",
        "The analysis identified 1,088.4 km of the 1,446.9 km KKH proxy above the map-wide P90 score. Of that high-score length, 295.6 km was supported and 792.8 km was classified as verification priority. KKH rankings remained stable after road-distance removal (Spearman rho = 0.981).",
    )

    add_heading(doc, "3. Main scientific findings now supported", 1)
    for item in (
        "AlphaEarth embeddings provide complementary spatial information, but the strongest and best-calibrated model combines them with physically interpretable conventional factors.",
        "The fusion gain is modest rather than universal: ROC-AUC and Brier improvements are spatially supported, whereas the PR-AUC gain over Conventional is not statistically resolved.",
        "High pooled performance does not imply uniform regional generalisation; Balochistan is the most difficult held-out transfer domain.",
        "Feature-space support and discrimination answer different questions. A domain can lie substantially inside AoA while still showing weaker classification performance, or show strong discrimination despite restricted AoA coverage.",
        "Control selection and failure typology materially influence apparent performance, but the fused model retains useful discrimination under more demanding matched-control experiments.",
        "Road proximity contributes strongly to absolute scores, yet most corridor-level ranking persists after its removal, indicating that the prioritisation pattern is not solely an accessibility artefact.",
    ):
        add_bullet(doc, item)

    add_heading(doc, "4. Why the current methodology is more publishable", 1)
    add_body(
        doc,
        "The manuscript no longer relies on a single accuracy statistic or a single random sample configuration. Each principal claim is paired with a diagnostic that tests its validity: spatial separation for predictive skill, paired block bootstrap for feature-set differences, matched controls for sampling bias, typology experiments for inventory composition, ablation for road dependence, LODO for regional transfer, AoA for predictor-domain support, and model disagreement for operational confidence.",
    )
    add_callout(
        doc,
        "Main novelty in simple terms",
        "We combine interpretable landslide-conditioning factors with AlphaEarth foundation-model embeddings, then test not only whether the model predicts well, but where it can transfer, where its environmental support is weak, how sensitive it is to the inventory and controls, and which CPEC-KKH road sections should be verified first.",
        fill=LIGHT_GREEN,
        accent=GREEN,
    )

    add_heading(doc, "5. Limits stated explicitly", 1)
    add_body(
        doc,
        "The manuscript deliberately limits its claims. The historical 1970-2020 inventory is related to a common baseline environmental representation; the resulting maps are spatial susceptibility screening products, not annual event probabilities. The 250 m grid is appropriate for corridor-scale prioritisation but does not replace slope-scale engineering assessment. AoA indicates environmental support, road exposure indicates inspection burden, and neither is treated as independent landslide validation.",
    )

    doc.core_properties.title = "Scientific and methodological improvements in the CPEC manuscript"
    doc.core_properties.subject = "Supervisor review brief"
    doc.core_properties.author = "Mohibullah"
    doc.core_properties.comments = "Generated from verified manuscript methods and results; no values were changed."
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
