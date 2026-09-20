"""Create a supervisor-facing Word brief summarising manuscript v3.2 improvements."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\DING PROJECT\FINAL PAPER\Manuscript_single_file\v3_2_2026-07-17")
OUTPUT = ROOT / "Teacher_brief_manuscript_v3_2_improvements_2026-07-17.docx"

BLUE = "1F4E79"
DARK_BLUE = "17365D"
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "666666"
GREEN = "2E7D5A"


def set_font(run, size=11, bold=False, color="000000", italic=False):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=60, start=110, bottom=60, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_font(run, size=9, color=MID_GRAY)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.10
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_font(lead, bold=True)
        rest = p.add_run(text[len(bold_lead):])
        set_font(rest)
    else:
        run = p.add_run(text)
        set_font(run)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.50)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.10
    run = p.add_run(text)
    set_font(run)
    return p


def add_callout(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.line_spacing = 1.10
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), LIGHT_BLUE)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), BLUE)
    borders.append(left)
    p_pr.append(borders)
    run = p.add_run(text)
    set_font(run, size=11, bold=True, color=DARK_BLUE)
    return p


def build() -> Path:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.70)
    section.bottom_margin = Inches(0.70)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for name, size, before, after, color in (
        ("Heading 1", 16, 16, 8, BLUE),
        ("Heading 2", 13, 12, 6, BLUE),
        ("Heading 3", 12, 8, 4, DARK_BLUE),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hrun = header.add_run("CPEC landslide susceptibility manuscript | Revision brief")
    set_font(hrun, size=9, bold=True, color=MID_GRAY)
    add_page_number(section.footer.paragraphs[0])

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(5)
    krun = kicker.add_run("SUPERVISOR REVIEW NOTE")
    set_font(krun, size=10, bold=True, color=GREEN)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    trun = title.add_run("Summary of improvements in manuscript v3.2")
    set_font(trun, size=23, bold=True, color="000000")

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    srun = subtitle.add_run("Transferability-aware landslide susceptibility mapping along the China-Pakistan Economic Corridor")
    set_font(srun, size=13, color=MID_GRAY)

    metadata = doc.add_paragraph()
    metadata.paragraph_format.space_after = Pt(8)
    metadata.paragraph_format.line_spacing = 1.05
    fields = [
        ("Prepared for", "Supervisor review"),
        ("Prepared by", "Mohibullah"),
        ("Version", "v3.2 relative to v3.1"),
        ("Date", date(2026, 7, 17).strftime("%d %B %Y")),
    ]
    for index, (label, value) in enumerate(fields):
        if index:
            separator = metadata.add_run("   |   ")
            set_font(separator, size=9.5, color=MID_GRAY)
        lr = metadata.add_run(f"{label}: ")
        set_font(lr, size=9.5, bold=True, color=DARK_BLUE)
        vr = metadata.add_run(value)
        set_font(vr, size=9.5)

    add_callout(
        doc,
        "The v3.2 revision improves scientific communication and figure readability without changing the datasets, model outputs, performance values, tables, or conclusions. The expanded Results introduced in v3.1 are retained in full.",
    )

    add_heading(doc, "1. Purpose of the revision", 1)
    add_body(
        doc,
        "The revision followed a complete visual audit of the manuscript figures. It removes overlaps and clipped labels, reduces redundant bar-chart use, and strengthens visual hierarchy at journal page size. All nine figures were regenerated from verified project data and checked individually and within the assembled Word manuscript.",
    )

    add_heading(doc, "2. What remained scientifically unchanged", 1)
    for item in (
        "The official CPEC boundary, inventory, 1,508 slope-failure samples, 1,808 controls, 19 conventional factors, 64 AlphaEarth dimensions, and three feature sets.",
        "The fully nested buffered spatial cross-validation design, six base learners, and L2 stacked ensemble.",
        "All ROC-AUC, PR-AUC, Brier score, calibration, TreeSHAP, robustness, LODO, AoA, and road-exposure values.",
        "The seven Results subsections and tables, Discussion, Conclusions, and overall scientific interpretation.",
    ):
        add_bullet(doc, item)

    add_heading(doc, "3. Main improvements from v3.1 to v3.2", 1)
    rows = [
        ("1", "Study-area map", "A large blank band separated global panel (a) from regional panel (b).", "Matched panel heights to geographic aspect ratios and anchored the locator maps together.", "Faster global-to-regional reading and more efficient use of page space."),
        ("2", "Baseline score maps", "The distribution key covered the score-density tails and the layout contained avoidable whitespace.", "Moved the key into an unused region and enlarged the mapped surfaces.", "The spatial comparison and distribution tails are now simultaneously readable."),
        ("3", "Robustness results", "The delta panel relied on another grouped bar chart.", "Replaced bars with a zero-centred multimetric dot display.", "Positive and negative departures from the primary fused model are easier to compare."),
        ("4", "Subdomain transfer", "The ROC-AUC/PR-AUC legend competed with labels and confidence intervals.", "Placed the shared legend outside the data region.", "All LODO estimates and 95% confidence intervals remain unobstructed."),
        ("5", "AoA comparison", "Grouped bars and their legend duplicated information already visible in the maps.", "Replaced the bars with an annotated domain-by-feature-set matrix on the same colour scale.", "Cross-domain support differences can be compared directly and precisely."),
        ("6", "Road prioritisation", "The KKH legend covered the map; the secondary axis title was clipped; two panels were bar-heavy.", "Moved map legends outside the data, used connected P80/P90 markers, and separated route length from no-road model scores in aligned dot plots.", "Road exposure, verification priority, and ablation stability are clearer without hiding mapped evidence."),
        ("7", "Whole figure set", "Panel spacing, legend placement, and visual encodings were not fully standardised.", "Applied consistent Arial typography, panel labels, colour meanings, map scales, line weights, and reserved legend space.", "The figures now read as one coherent publication set rather than independent outputs."),
    ]
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["No.", "Area", "Previous limitation", "v3.2 improvement", "Scientific communication benefit"]
    for cell, text in zip(table.rows[0].cells, headers):
        shade_cell(cell, BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        set_font(run, size=9, bold=True, color="FFFFFF")
    for row_data in rows:
        cells = table.add_row().cells
        for col, (cell, text) in enumerate(zip(cells, row_data)):
            if len(table.rows) % 2 == 1:
                shade_cell(cell, LIGHT_GRAY)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col == 0 else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(text)
            set_font(run, size=8.3, bold=(col == 1), color="000000")
    set_table_geometry(table, [500, 1250, 2400, 2800, 2290])

    add_heading(doc, "4. Improvements to publication readiness", 1)
    add_body(doc, "The revision strengthens the manuscript in four reviewer-facing ways:")
    for item in (
        "Reviewer safety: no legend, panel label, colour bar, map element, or axis title now hides scientific information.",
        "Interpretive efficiency: maps remain the primary evidence, while interval, matrix, and dot displays replace avoidable bars.",
        "Cross-panel consistency: comparable feature sets and metrics retain common colours, scales, symbols, and terminology.",
        "Reproducibility: each figure is generated from explicit project inputs and exported as both a 500 dpi PNG and vector PDF.",
    ):
        add_bullet(doc, item)

    add_heading(doc, "5. Quality assurance completed", 1)
    for item in (
        "All nine standalone figures were inspected at full raster resolution.",
        "The complete manuscript was rendered through Microsoft Word and inspected across 28 pages.",
        "Seven tables and nine figures are numbered consecutively and match their in-text references and captions.",
        "The manuscript consistency audit passed all structural and stale-content checks.",
        "No generative-AI-created or AI-altered scientific figure was used.",
    ):
        add_bullet(doc, item)

    add_heading(doc, "6. Files supplied for review", 1)
    add_body(doc, "The review package contains the following primary materials:")
    for item in (
        "CPEC_manuscript_v3_2_figure_refined_2026-07-17.docx - complete revised manuscript.",
        "figures_png - nine high-resolution raster figures for inspection and submission use.",
        "figures_pdf - nine vector figures for publication-quality scaling.",
        "Teacher_brief_manuscript_v3_2_improvements_2026-07-17.docx - this revision summary.",
    ):
        add_bullet(doc, item)

    add_callout(
        doc,
        "Recommended review focus: the revised figure hierarchy, the map-led presentation of spatial results, the direct AoA comparison, the unobstructed LODO confidence intervals, and the clearer CPEC/KKH infrastructure-prioritisation outputs.",
    )

    core = doc.core_properties
    core.title = "Summary of improvements in manuscript v3.2"
    core.subject = "Supervisor-facing revision brief for the CPEC landslide susceptibility manuscript"
    core.author = "Mohibullah"
    core.keywords = "CPEC; landslide susceptibility; manuscript revision; figure quality assurance"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
